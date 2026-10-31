"""根据各学科错题集/题库.json生成两份可打印Word：
1) 空白练习版；2) 参考答案与讲解版。"""
from __future__ import annotations
import argparse, datetime as dt, json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUBJECT_FOLDERS = {
    "语文": "01_语文", "数学": "02_数学", "英语": "03_英语",
    "地理": "04_地理", "历史": "05_历史", "生物": "06_生物", "道法": "07_道法",
}
FONT = "宋体"
HEI = "黑体"

def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def add_run(p, text, name=FONT, size=11, bold=False, color=None):
    r = p.add_run(str(text))
    set_font(r, name, size, bold, color)
    return r

def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "第 ", size=8.5, color=(128,128,128))
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
    p._p.append(fld)
    add_run(p, " 页", size=8.5, color=(128,128,128))

def setup_doc(subject, variant):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT; st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.line_spacing = 1.3
    st.paragraph_format.space_after = Pt(4)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(hp, f"朱秋宇 · {subject}错题集 · {variant}", size=8.5, color=(128,128,128))
    add_page_number(sec)
    return doc

def add_title(doc, subject, variant, count):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    add_run(p, f"{subject}错题练习 · {variant}", HEI, 20, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, f"朱秋宇 · 七年级  |  共 {count} 题  |  生成日期：{dt.date.today():%Y-%m-%d}", size=9, color=(112,112,112))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(16)
    add_run(p, "姓名：__________    日期：__________    用时：________分钟    自我评分：________", size=10)

def add_meta(doc, q, n):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    add_run(p, f"第 {n} 题", HEI, 14, True)
    bits = []
    for label, key in (("来源", "source"), ("知识点", "topic"), ("题型", "type")):
        if q.get(key): bits.append(f"{label}：{q[key]}")
    if bits:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        add_run(p, "  |  ".join(bits), size=8.5, color=(112,112,112))

def add_text(doc, text, size=11, indent=False):
    for line in str(text or "").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        if indent: p.paragraph_format.first_line_indent = Cm(0.74)
        add_run(p, line if line.strip() else " ", size=size)

def add_image(doc, base, rel, max_cm=15):
    if not rel: return
    path = (base / rel).resolve()
    if not path.exists():
        add_text(doc, f"[图片缺失：{rel}]", 9); return
    try:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(max_cm))
    except Exception:
        add_text(doc, f"[图片无法插入：{rel}]", 9)

def add_blank_lines(doc, n):
    n = max(3, min(int(n or 8), 30))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
    add_run(p, "答题区（请写出过程）", HEI, 9.5, True, (64,64,64))
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.8
        add_run(p, "_" * 72, size=10, color=(205,205,205))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
    add_run(p, "我的答案：____________________________________________________________", size=10.5)

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def resolve_subject(arg):
    if not arg:
        return list(SUBJECT_FOLDERS.keys())
    arg = arg.strip()
    if arg in SUBJECT_FOLDERS:
        return [arg]
    for subj, folder in SUBJECT_FOLDERS.items():
        if arg == folder:
            return [subj]
    raise SystemExit(f"未找到学科：{arg}。可用值：{', '.join(SUBJECT_FOLDERS)} 或目录名")

def normalize_question(q):
    q = dict(q)
    q.setdefault("id", "")
    q.setdefault("date", "")
    q.setdefault("source", "")
    q.setdefault("type", "")
    q.setdefault("topic", "")
    q.setdefault("question", "")
    q.setdefault("question_display", "")
    q.setdefault("question_images", [])
    q.setdefault("blank_lines", 8)
    q.setdefault("answer", "")
    q.setdefault("explanation", "")
    q.setdefault("mistake_reason", "")
    q.setdefault("variant", "")
    q.setdefault("status", "待复习")
    if isinstance(q["question_images"], str):
        q["question_images"] = [q["question_images"]]
    return q

def build_blank(doc, subject, questions, folder, out_dir):
    add_title(doc, subject, "空白练习版", len(questions))
    for i, q in enumerate(questions, 1):
        if i > 1:
            page_break(doc)
        add_meta(doc, q, i)
        add_text(doc, q.get("question_display") or q["question"], size=11)
        for img in q["question_images"]:
            add_image(doc, folder / "错题集", img, max_cm=16)
        add_blank_lines(doc, q["blank_lines"])
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
        add_run(p, "这次我的错因（可做完后写）：______________________________________", size=9, color=(128,128,128))
    path = out_dir / f"{subject}_错题练习_空白版.docx"
    doc.save(path)
    return path

def build_answers(doc, subject, questions, folder, out_dir):
    add_title(doc, subject, "参考答案与讲解", len(questions))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
    add_run(p, "答案速查", HEI, 14, True)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "题号"; hdr[1].text = "答案 / 结论"
    for c in hdr:
        for r in c.paragraphs[0].runs:
            set_font(r, HEI, 10, True)
    for i, q in enumerate(questions, 1):
        row = t.add_row().cells
        row[0].text = f"第 {i} 题"
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[1].text = (str(q["answer"]).strip() or "❓待补")
        for cell in row:
            for r in cell.paragraphs[0].runs:
                set_font(r, FONT, 9.5)
    for i, q in enumerate(questions, 1):
        if i > 1:
            page_break(doc)
        add_meta(doc, q, i)
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
        add_run(p, "题目回顾：", HEI, 9.5, True)
        add_text(doc, q.get("question_display") or q["question"], size=9.5, indent=True)
        for img in q["question_images"]:
            add_image(doc, folder / "错题集", img, max_cm=15)
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
        add_run(p, "答案：", HEI, 11, True)
        add_run(p, (str(q["answer"]).strip() or "❓待补"), size=11, bold=True, color=(192,0,0))
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
        add_run(p, "解析：", HEI, 11, True)
        add_text(doc, q["explanation"] or "❓待补", size=10.5, indent=True)
        if q["mistake_reason"]:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
            add_run(p, "错因提醒 / 复盘重点：", HEI, 10.5, True)
            add_text(doc, q["mistake_reason"], size=10, indent=True)
        if q["variant"]:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
            add_run(p, "同类变式：", HEI, 10.5, True)
            add_text(doc, q["variant"], size=10, indent=True)
        if q["status"]:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5)
            add_run(p, f"掌握状态：{q['status']}", size=8.5, color=(128,128,128))
    path = out_dir / f"{subject}_错题练习_参考答案与讲解.docx"
    doc.save(path)
    return path

def build_subject(subject):
    folder = ROOT / SUBJECT_FOLDERS[subject]
    bank_path = folder / "错题集" / "题库.json"
    if not bank_path.exists():
        print(f"[跳过] {subject}: 未找到 {bank_path.relative_to(ROOT)}")
        return None
    try:
        data = json.loads(bank_path.read_text(encoding="utf-8"))
    except Exception as e:
        raise SystemExit(f"[错误] {subject} 题库 JSON 解析失败：{e}")
    questions = [normalize_question(q) for q in (data.get("questions") or [])]
    if not questions:
        print(f"[跳过] {subject}: 题库为空")
        return None
    out_dir = folder / "错题集"
    out_dir.mkdir(parents=True, exist_ok=True)
    blank = build_blank(setup_doc(subject, "空白练习版"), subject, questions, folder, out_dir)
    ans = build_answers(setup_doc(subject, "参考答案与讲解"), subject, questions, folder, out_dir)
    print(f"[生成] {subject}: {blank.relative_to(ROOT)} | {ans.relative_to(ROOT)} | {len(questions)} 题")
    return blank, ans, len(questions)

def main():
    parser = argparse.ArgumentParser(description="生成错题空白练习版和参考答案版 Word")
    parser.add_argument("subject", nargs="?", help="学科名或目录名；不填则生成全部已有题库")
    parser.add_argument("--all", action="store_true", help="生成所有已有题库的学科")
    args = parser.parse_args()
    subjects = list(SUBJECT_FOLDERS.keys()) if args.all else resolve_subject(args.subject)
    built = 0
    for s in subjects:
        if build_subject(s):
            built += 1
    print(f"完成：生成 {built} 个学科")
    if built > 0:
        sync_script = ROOT / "tools" / "sync_to_deeptutor.py"
        dt_python = Path.home() / ".deeptutor-venv" / "bin" / "python"
        py_bin = str(dt_python) if dt_python.exists() else sys.executable
        if sync_script.exists():
            print("[自动同步] 正在自动将最新错题录入 DeepTutor 知识库...")
            import subprocess
            res = subprocess.run([py_bin, str(sync_script), "--wrongbook"], capture_output=True, text=True)
            if res.returncode == 0:
                print("  ✅ DeepTutor 错题知识库自动同步完成！")
            else:
                print(f"  ⚠️ DeepTutor 自动同步提示: {res.stderr.strip() or res.stdout.strip()}")

if __name__ == "__main__":
    main()
